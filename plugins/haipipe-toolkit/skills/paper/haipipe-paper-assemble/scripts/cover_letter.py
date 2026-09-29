"""run-delivery-coverletter: the submission cover letter, built like the manuscript (0.9.0, JL 260929).

WORDS  the "Cover letter" division of the submission Round page
       (B<x>-<desk>-Round/RD<NN>-<desk>-submission-<YYYYMMDD>/<stem>.md). It is approved like any
       page. While the page Content has no such division, the Round page's latest draft is used and
       the letter stays not ready.
FACTS  filled by code so they never drift from the manuscript: {title} {journal} {article_type}
       {pages} {tables} {figures} in the words, plus the date, addressee and author block from
       paper-build.toml [coverletter].
OUTPUT delivery/cover-letter/ (the [outputs] cover_letter_pdf / cover_letter_docx paths): a PDF
       compiled from generated LaTeX and a DOCX, both generated, never edited by hand.
CHECKS every number in the letter is printed in the manuscript · required mentions present (profile
       cover_letter_must_mention + paper must_mention) · no causal verbs · no process text ·
       manuscript pages within the venue cap · author block complete · the letter fits two pages.
"""
from __future__ import annotations

import datetime
import re
import shutil
import subprocess
from pathlib import Path

HEAD = re.compile(r"(?m)^(#{2,6})\s+(.*)$")
FACT = re.compile(r"\{(title|journal|article_type|pages|tables|figures)\}")
NUMBER = re.compile(r"(?<![\w.])\d[\d,]*(?:\.\d+)?")
CAUSAL = re.compile(r"(?i)\b(causes?|caused|causal effect|leads? to|led to|drives?|driven by|results? in|impact of|effect of)\b")
PROCESS = re.compile(r"\[(?:TODO|TOADD|Q-|AUTHOR|confirm)[^\]]*\]|\bTOADD\b|\bTODO\b", re.I)


def division(text: str, name: str = "cover letter") -> list[str] | None:
    """Paragraphs of the first heading whose title contains `name`, up to the next heading of the
    same or a higher level. Headings, tables, lists, quotes and comments inside it are not letter text."""
    heads = list(HEAD.finditer(text))
    for i, h in enumerate(heads):
        if name not in h.group(2).lower():
            continue
        level, end = len(h.group(1)), len(text)
        for nxt in heads[i + 1:]:
            if len(nxt.group(1)) <= level:
                end = nxt.start()
                break
        body = re.sub(r"<!--.*?-->", "", text[h.end():end], flags=re.S)
        paras = []
        for block in re.split(r"\n\s*\n", body):
            lines = [l for l in block.strip().splitlines() if l.strip()]
            if not lines or lines[0].lstrip()[:1] in "#|>-*" and not lines[0].lstrip().startswith("*"):
                continue
            paras.append(" ".join(l.strip() for l in lines))
        return [p for p in paras if p]
    return None


def round_page(root: Path, rd: str) -> Path | None:
    hits = [p for p in root.glob(f"B*-*-Round/{rd}-*/{rd}-*.md") if p.parent.name == p.stem]
    return hits[0] if len(hits) == 1 else None


def _draft_version(p: Path):
    m = re.search(r"-draft-v(\d+)(?:\.(\d+))?\.md$", p.name)
    return (int(m.group(1)), int(m.group(2) or 0)) if m else (-1, -1)


def letter_source(page: Path):
    """(paragraphs, where, path): Content first; else the page's latest draft."""
    words = division(page.read_text(encoding="utf-8"))
    if words:
        return words, "content", page
    drafts = sorted(page.parent.glob("draft/*-draft-v*.md"), key=_draft_version)
    for d in reversed(drafts):
        words = division(d.read_text(encoding="utf-8"))
        if words:
            return words, "draft", d
    return [], "missing", page


def pdf_pages(pdf: Path) -> int | None:
    if not pdf.exists():
        return None
    try:
        out = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True).stdout
        m = re.search(r"(?m)^Pages:\s+(\d+)", out)
        if m:
            return int(m.group(1))
    except FileNotFoundError:
        pass
    return len(re.findall(rb"/Type\s*/Page[^s]", pdf.read_bytes())) or None


def pdf_text(pdf: Path) -> str | None:
    try:
        r = subprocess.run(["pdftotext", str(pdf), "-"], capture_output=True, text=True)
        return r.stdout if r.returncode == 0 else None
    except FileNotFoundError:
        return None


def fill(text: str, facts: dict) -> str:
    return FACT.sub(lambda m: str(facts.get(m.group(1), m.group(0))), text)


# ── rendering ─────────────────────────────────────────────────────────────────
_TEX = {"\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#", "_": r"\_",
        "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}", "^": r"\textasciicircum{}"}


def tex_escape(s: str) -> str:
    out = "".join(_TEX.get(c, c) for c in s)
    out = re.sub(r"\*([^*]+)\*", r"\\emph{\1}", out)
    out = re.sub(r'"([^"]*)"', r"``\1''", out)          # straight quotes print as two closing quotes in TeX
    return out


def write_tex(path: Path, letter: dict) -> None:
    sign = [tex_escape(x) for x in (letter["signatory"], letter["affiliation"], letter["email"]) if x]
    body = "\n\n".join(tex_escape(p) for p in letter["paragraphs"])
    authors = letter["on_behalf"]
    path.write_text(
        "% GENERATED by haipipe-paper-assemble cover_letter.py; edit the Round page, then rebuild\n"
        "\\documentclass[12pt]{article}\n\\usepackage[letterpaper,margin=1in]{geometry}\n\\usepackage{fontspec}\n"
        "\\IfFontExistsTF{TeX Gyre Termes}{\\setmainfont{TeX Gyre Termes}}{\\IfFontExistsTF{Times New Roman}{\\setmainfont{Times New Roman}}{}}\n"
        "\\usepackage[hidelinks]{hyperref}\n\\setlength{\\parindent}{0pt}\n\\setlength{\\parskip}{0.8em}\n\\pagestyle{empty}\n"
        "\\begin{document}\n"
        f"{tex_escape(letter['date'])}\n\n"
        f"{tex_escape(letter['addressee'])}\\\\\n{tex_escape(letter['journal'])}\n\n"
        f"Dear {tex_escape(letter['salutation'])},\n\n{body}\n\n"
        + (f"{tex_escape(authors)}\n\n" if authors else "")
        + "Sincerely,\n\n\\vspace{1.5em}\n" + "\\\\\n".join(sign) + "\n\\end{document}\n",
        encoding="utf-8")


def write_docx(path: Path, letter: dict) -> None:
    from docx import Document
    from docx.shared import Inches, Pt

    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(1)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    rpr = normal.element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "Times New Roman")

    def para(text: str, after: float = 10):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.line_spacing = 1.0
        for i, chunk in enumerate(re.split(r"\*([^*]+)\*", text)):
            if chunk:
                p.add_run(chunk).italic = bool(i % 2)
        return p

    para(letter["date"], 14)
    para(letter["addressee"], 0)
    para(letter["journal"], 14)
    para(f"Dear {letter['salutation']},")
    for p in letter["paragraphs"]:
        para(p)
    if letter["on_behalf"]:
        para(letter["on_behalf"])
    para("Sincerely,", 24)
    for i, line in enumerate(x for x in (letter["signatory"], letter["affiliation"], letter["email"]) if x):
        para(line, 0)
    doc.save(path)


# ── the lane ──────────────────────────────────────────────────────────────────
def build_cover_letter(cfg: dict, profile: dict, root: Path, here: Path, out: dict, *,
                       title: str, main_pdf: Path, tables: int, figures: int, run) -> dict | None:
    """Build the letter when paper-build.toml has [coverletter]; return the manifest report."""
    cl = cfg.get("coverletter")
    if not isinstance(cl, dict):
        return None
    pdf_out = (here / out["cover_letter_pdf"]).resolve() if out.get("cover_letter_pdf") else None
    docx_out = (here / out["cover_letter_docx"]).resolve() if out.get("cover_letter_docx") else None
    if pdf_out is None and docx_out is None:
        return {"ready": False, "checks": [], "missing": ["[outputs] cover_letter_pdf / cover_letter_docx"]}

    rd = str(cl.get("round", ""))
    page = round_page(root, rd) if rd else None
    paragraphs, where, source = letter_source(page) if page else ([], "missing", None)
    pages = pdf_pages(main_pdf)
    facts = {"title": title, "journal": cl.get("journal", profile.get("venue_label", "")),
             "article_type": cl.get("article_type", ""), "pages": pages or "", "tables": tables, "figures": figures}
    filled = [fill(p, facts) for p in paragraphs]
    authors = [a for a in cl.get("authors", []) if a]
    signatory = cl.get("signatory", authors[0] if authors else "")
    others = [a for a in authors if a != signatory]
    letter = {
        "date": (lambda d: f"{d:%B} {d.day}, {d.year}")(datetime.date.today()),
        "addressee": cl.get("addressee", "Editor-in-Chief"), "journal": facts["journal"],
        "salutation": cl.get("salutation", cl.get("addressee", "Editor-in-Chief")),
        "paragraphs": filled, "signatory": signatory,
        "affiliation": cl.get("affiliation", ""), "email": cl.get("email", ""),
        "on_behalf": (f"I submit this manuscript on behalf of my co-authors, {', '.join(others[:-1]) + ' and ' if len(others) > 1 else ''}{others[-1]}."
                      if others else ""),
    }

    checks = []
    def check(name, ok, detail=""):
        checks.append({"check": name, "ok": bool(ok), "detail": detail})

    check("source", where == "content",
          {"content": "the Round page's Content", "draft": f"the draft {source.name if source else ''} (not yet adopted)",
           "missing": f"no Cover letter division on the {rd} Round page" if page else f"no Round page {rd}"}[where])
    text = " ".join(filled)
    body_only = " ".join(FACT.sub("", p) for p in paragraphs)       # facts filled by code are true by construction
    manuscript = pdf_text(main_pdf) or ""
    if manuscript:
        flat = re.sub(r"\s+", " ", manuscript)
        absent = sorted({n for n in NUMBER.findall(body_only) if n.rstrip(".,") not in flat})
        check("numbers match the manuscript", not absent, ", ".join(absent) or "every number is printed in the manuscript")
    else:
        check("numbers match the manuscript", False, "no manuscript text (pdftotext missing or no PDF)")
    musts = list(profile.get("cover_letter_must_mention", [])) + list(cl.get("must_mention", []))
    lacking = [m for m in musts if m.lower() not in (text + " " + letter["journal"]).lower()]
    check("required mentions", not lacking, ", ".join(lacking) or "all present")
    causal = sorted({m.group(0).lower() for m in CAUSAL.finditer(text)})
    check("associational language", not causal, ", ".join(causal) or "no causal verbs")
    process = sorted({m.group(0) for m in PROCESS.finditer(text)})
    check("no process text", not process, ", ".join(process) or "clean")
    cap = profile.get("page_cap")
    if cap:
        check("manuscript within the page cap", pages is not None and pages <= int(cap), f"{pages} of {cap} pages")
    missing = [k for k in ("signatory", "affiliation", "email") if not letter[k]]
    check("author block complete", not missing, ", ".join(missing) or "complete")

    rc = None
    if pdf_out:
        pdf_out.parent.mkdir(parents=True, exist_ok=True)
        tex = pdf_out.with_suffix(".tex")
        write_tex(tex, letter)
        r = run(["latexmk", "-xelatex", "-interaction=nonstopmode", "-halt-on-error", "-quiet", tex.name],
                cwd=tex.parent, capture_output=True, text=True)
        rc = r.returncode
        for junk in tex.parent.glob(tex.stem + ".*"):
            if junk.suffix not in {".tex", ".pdf", ".log"}:
                junk.unlink()
        n = pdf_pages(pdf_out)
        check("letter fits two pages", n is not None and n <= 2, f"{n} page(s)")
    if docx_out:
        docx_out.parent.mkdir(parents=True, exist_ok=True)
        write_docx(docx_out, letter)

    return {
        "round": rd, "source": str(source.relative_to(root)) if source else None, "from": where,
        "ready": all(c["ok"] for c in checks) and rc in (0, None),
        "latexmk_rc": rc, "checks": checks,
        "outputs": {"pdf": out.get("cover_letter_pdf") if pdf_out and pdf_out.exists() else None,
                    "docx": out.get("cover_letter_docx") if docx_out and docx_out.exists() else None},
    }
