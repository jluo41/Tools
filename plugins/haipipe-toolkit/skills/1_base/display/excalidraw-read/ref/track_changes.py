"""A copy of a Word file with tracked changes: each edit is the smallest Word deletion and insertion by one author.

excalidraw-read's revise mode writes the delivery with it (read_revise.py deliver); first written for b04_prep_ai's
delivery/ (2026-10-11), shared here so every Block builds its tracked file the same way.

    from track_changes import write_tracked, preview
    write_tracked(template, edits, out, author="JL", figure=None)   # the Word file
    preview(template, edits)                                         # what each edit changes, for a drawing

An edit is a dict (s04's yaml):
  find      the template's exact words, inside one paragraph or table-cell line. Found once in the file, or,
            with `after`, the first match after each `after` phrase in turn ([Year 1, "Q3:"]).
  new       what they become; only the words that differ are marked (Word shows the deletion struck, the
            insertion underlined). "" deletes the find; no `new` changes nothing and only adds comments.
            A line break in `new` starts a new paragraph in the style of the one it sits in; it is allowed
            only where the change ends its paragraph.
  image     "figure" (or a path): the find's table cell, else its paragraph, gives way to that picture.
  split     instead of find: a phrase that should start a new paragraph. Its paragraph is cut before it, and
            the new paragraph mark is a tracked insertion (rejecting it joins the two again). A phrase already
            at a paragraph's start changes nothing.
  evidence  {need, ask, who, anchor, found?}, or a list of them: one Word comment each, "Evidence needed: ...",
            on its anchor words, looked for first in the changed paragraphs, then in the whole file.
A placeholder in brackets ("[Your text here]") hands its place to plain text: its grey and italic are dropped.
Every package part other than the document, its relationships, comments and content types keeps the
template's bytes. The build stops when a find or an anchor is missing, or a find is not unique.
"""
import datetime
import re
import tempfile
import zipfile
from copy import deepcopy
from pathlib import Path

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from docx.text.run import Run

W = qn
EDITED = {"word/document.xml", "word/_rels/document.xml.rels", "[Content_Types].xml", "word/comments.xml"}
PLACEHOLDER = re.compile(r"\s*\[[^\]]*\]\s*")


# ── where the words are ──────────────────────────────────────────────────────────────────────────────

def paragraphs(body):
    """Every paragraph of the body in reading order, table cells included (not text boxes)."""
    return [p for p in body.iter(W("w:p")) if not any(a.tag == W("w:txbxContent") for a in p.iterancestors())]


def runs_of(p):
    """The paragraph's own text runs in order, those inside insertions included, deleted ones left out."""
    out = []
    for r in p.iter(W("w:r")):
        up = r.getparent()
        while up is not p and up.tag not in (W("w:p"), W("w:del")):
            up = up.getparent()
        if up is p:
            out.append(r)
    return out


def run_text(r):
    return "".join(t.text or "" for t in r.findall(W("w:t")))


def para_text(p):
    return "".join(run_text(r) for r in runs_of(p))


def locate(paras, edit):
    """(paragraph index, start, end) of the edit's find, per the rules above."""
    texts = [para_text(p) for p in paras]
    stream = "\n".join(texts)
    starts, k = [], 0
    for t in texts:
        starts.append(k)
        k += len(t) + 1
    pos = 0
    for phrase in edit.get("after", []) or []:
        i = stream.find(str(phrase), pos)
        if i < 0:
            raise SystemExit(f"edit {edit['id']}: after-phrase {phrase!r} not found")
        pos = i + len(str(phrase))
    find = edit["find"]
    i = stream.find(find, pos)
    if i < 0:
        raise SystemExit(f"edit {edit['id']}: find {find!r} not found")
    if not edit.get("after") and stream.count(find) != 1:
        raise SystemExit(f"edit {edit['id']}: find {find!r} appears {stream.count(find)} times; give `after`")
    if "\n" in stream[i:i + len(find)]:
        raise SystemExit(f"edit {edit['id']}: find {find!r} crosses a paragraph")
    n = max(j for j, s in enumerate(starts) if s <= i)
    return n, i - starts[n], i - starts[n] + len(find)


def split_change(old, new):
    """The smallest change on word edges: (kept before, removed, added, kept after)."""
    a = 0
    while a < min(len(old), len(new)) and old[a] == new[a]:
        a += 1
    b = 0
    while b < min(len(old), len(new)) - a and old[len(old) - 1 - b] == new[len(new) - 1 - b]:
        b += 1

    def inside(s, k):
        return 0 < k < len(s) and s[k - 1].isalnum() and s[k].isalnum()
    while a > 0 and (inside(old, a) or inside(new, a)):
        a -= 1
    while b > 0 and (inside(old, len(old) - b) or inside(new, len(new) - b)):
        b -= 1
    return old[:a], old[a:len(old) - b], new[a:len(new) - b], old[len(old) - b:]


# ── preview: what the drawing shows ──────────────────────────────────────────────────────────────────

def sentence_around(text, start, end):
    """The sentence(s) holding text[start:end]."""
    s = max(text.rfind(". ", 0, start), text.rfind(".  ", 0, start))
    s = 0 if s < 0 else s + 2
    e = text.find(". ", max(start, end - 1))           # a find that ends its sentence keeps it
    e = len(text) if e < 0 else e + 1
    return s, e


def preview(template, edits):
    """For each edit: the template's sentence split into kept, removed and kept, and the new words in its place."""
    doc = docx.Document(template)
    paras = paragraphs(doc.element.body)
    out = []
    for edit in edits:
        n, i, j = locate(paras, edit)
        text = para_text(paras[n])
        s, e = sentence_around(text, i, j)
        new = edit.get("new")
        before, removed, added, after = split_change(edit["find"], edit["find"] if new is None else new)
        if edit.get("image"):
            before, removed, added, after = "", edit["find"], "", ""
        out.append({"edit": edit, "before": text[s:i] + before, "removed": removed, "added": added,
                    "after": after + text[j:e], "in_table": any(a.tag == W("w:tc") for a in paras[n].iterancestors())})
    return out


# ── writing the tracked changes ──────────────────────────────────────────────────────────────────────

class Tracker:
    def __init__(self, doc, author):
        self.doc, self.author, self.next_id = doc, author, 9000
        self.date = datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")

    def mark(self, tag):
        e = OxmlElement(tag)
        e.set(W("w:id"), str(self.next_id))
        e.set(W("w:author"), self.author)
        e.set(W("w:date"), self.date)
        self.next_id += 1
        return e

    @staticmethod
    def set_text(r, text):
        for t in r.findall(W("w:t")):
            r.remove(t)
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = text
        r.append(t)

    def cut(self, p, k):
        """A run edge at character k of the paragraph."""
        pos = 0
        for r in runs_of(p):
            t = run_text(r)
            if pos < k < pos + len(t):
                r2 = deepcopy(r)
                self.set_text(r, t[:k - pos])
                self.set_text(r2, t[k - pos:])
                r.addnext(r2)
                return
            pos += len(t)

    def runs_between(self, p, i, j):
        self.cut(p, i)
        self.cut(p, j)
        out, pos = [], 0
        for r in runs_of(p):
            t = run_text(r)
            if pos >= i and pos + len(t) <= j and t:
                out.append(r)
            pos += len(t)
        return out

    def delete(self, runs):
        for r in runs:
            d = self.mark("w:del")
            r.addprevious(d)
            d.append(r)
            for t in r.findall(W("w:t")):
                t.tag = W("w:delText")
        return runs[-1].getparent() if runs else None

    def new_run(self, like, text, plain):
        r = OxmlElement("w:r")
        if like is not None and like.find(W("w:rPr")) is not None:
            rpr = deepcopy(like.find(W("w:rPr")))
            if plain:
                for tag in ("w:color", "w:i", "w:iCs", "w:highlight"):
                    for c in rpr.findall(W(tag)):
                        rpr.remove(c)
            r.append(rpr)
        self.set_text(r, text)
        return r

    def insert_after(self, anchor_el, p, run):
        ins = self.mark("w:ins")
        ins.append(run)
        if anchor_el is None:                       # at the paragraph's start
            ppr = p.find(W("w:pPr"))
            (ppr.addnext if ppr is not None else lambda e: p.insert(0, e))(ins)
        else:
            anchor_el.addnext(ins)
        return ins

    def split(self, p, k):
        """Cut paragraph p before character k: p keeps the text before, a new paragraph after it takes the rest
        and p's own paragraph mark; p's new mark is a tracked insertion. Returns the new paragraph."""
        self.cut(p, k)
        pos, first = 0, None
        for r in runs_of(p):
            t = run_text(r)
            if pos >= k and t:
                first = r
                break
            pos += len(t)
        if first is None:
            raise SystemExit(f"split: no text after character {k}")
        top = first
        while top.getparent() is not p:              # inside an insertion: cut the insertion in two
            up = top.getparent()
            if top.getprevious() is not None:
                twin = deepcopy(up)
                for c in list(twin):
                    twin.remove(c)
                twin.set(W("w:id"), str(self.next_id))
                self.next_id += 1
                move, nxt = [], top
                while nxt is not None:
                    move.append(nxt)
                    nxt = nxt.getnext()
                for c in move:
                    twin.append(c)
                up.addnext(twin)
                top = twin
            else:
                top = up
        p2 = OxmlElement("w:p")
        ppr = p.find(W("w:pPr"))
        if ppr is not None:
            p2.append(deepcopy(ppr))                 # the second part keeps the paragraph's own mark
        move, nxt = [], top
        while nxt is not None:
            move.append(nxt)
            nxt = nxt.getnext()
        for c in move:
            p2.append(c)
        p.addnext(p2)
        if ppr is None:
            ppr = OxmlElement("w:pPr")
            p.insert(0, ppr)
        rpr = ppr.find(W("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            ppr.append(rpr)
        rpr.insert(0, self.mark("w:ins"))            # the first part's mark is the new one
        return p2

    def new_paragraph(self, after_p, like_run, text, plain):
        p = OxmlElement("w:p")
        ppr = deepcopy(after_p.find(W("w:pPr"))) if after_p.find(W("w:pPr")) is not None else OxmlElement("w:pPr")
        rpr = ppr.find(W("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            ppr.append(rpr)
        rpr.insert(0, self.mark("w:ins"))          # the paragraph mark itself is new (w:ins comes first)
        p.append(ppr)
        ins = self.mark("w:ins")
        ins.append(self.new_run(like_run, text, plain))
        p.append(ins)
        after_p.addnext(p)
        return p


def write_tracked(template, edits, out, author="JL", figure=None, image_width=6.2):
    """Write `out`: the template with every edit as a tracked change and every evidence item as a comment."""
    template, out = Path(template), Path(out)
    before_bytes = template.read_bytes()
    doc = docx.Document(template)
    body = doc.element.body
    tr = Tracker(doc, author)
    touched = {}                                    # edit id -> the paragraphs it changed or added
    for edit in edits:
        paras = paragraphs(body)
        if edit.get("split"):
            n, i, _ = locate(paras, {"id": edit["id"], "find": edit["split"]})
            p = paras[n]
            touched[edit["id"]] = [p] + ([tr.split(p, i)] if para_text(p)[:i].strip() else [])
            continue
        n, i, j = locate(paras, edit)
        p = paras[n]
        here = [p]
        if edit.get("image"):
            cell = next((a for a in p.iterancestors() if a.tag == W("w:tc")), None)
            for q in ([x for x in cell.iter(W("w:p"))] if cell is not None else [p]):
                tr.delete(runs_of(q))
            path = Path(figure if edit["image"] == "figure" else edit["image"])
            run = Paragraph(p, doc.part).add_run()
            run.add_picture(str(path), width=Inches(image_width))
            p.remove(run._r)
            tr.insert_after(p[-1] if len(p) else None, p, run._r)
        elif edit.get("new") is not None:
            text = para_text(p)
            keep_a, removed, added, keep_b = split_change(edit["find"], edit["new"])
            a, b = i + len(keep_a), j - len(keep_b)
            if "\n" in added and text[b:].strip():
                raise SystemExit(f"edit {edit['id']}: a line break is allowed only where the change ends its paragraph")
            gone = tr.runs_between(p, a, b) if removed else []
            all_runs = runs_of(p)
            if gone:
                like = gone[0]
            else:
                tr.cut(p, a)
                pos, like = 0, None
                for r in runs_of(p):
                    if pos + len(run_text(r)) <= a:
                        like = r
                    pos += len(run_text(r))
                like = like if like is not None else (all_runs[0] if all_runs else None)
            plain = bool(removed) and PLACEHOLDER.fullmatch(removed) is not None
            at = tr.delete(gone) if gone else None
            if at is None:                           # the run just before the insertion point
                pos, at = 0, None
                for r in runs_of(p):
                    if pos + len(run_text(r)) <= a:
                        at = r if r.getparent() is p else r.getparent()
                    pos += len(run_text(r))
            lines = added.split("\n")
            if lines[0]:
                tr.insert_after(at, p, tr.new_run(like, lines[0], plain))
            last = p
            for line in lines[1:]:
                last = tr.new_paragraph(last, like, line, plain)
                here.append(last)
        touched[edit["id"]] = here

    n_comments = 0
    for edit in edits:
        items = edit.get("evidence") or []
        for ev in (items if isinstance(items, list) else [items]):
            runs = None
            for scope in (touched[edit["id"]], paragraphs(body)):
                for q in scope:
                    k = para_text(q).find(ev["anchor"])
                    if k >= 0:
                        runs = [Run(r, Paragraph(q, doc.part)) for r in tr.runs_between(q, k, k + len(ev["anchor"]))]
                        break
                if runs:
                    break
            if not runs:
                raise SystemExit(f"edit {edit['id']}: anchor {ev['anchor']!r} not found")
            who = ", ".join(ev.get("who", []))
            ask = ev.get("ask", "")
            found = "; ".join(f["cite"] for f in ev.get("found") or [])      # papers a Discovery Run found
            text = (f"Found: {found}. " if found else "") + f"Evidence needed: {ev['need']}." \
                + (f" {ask[:1].upper() + ask[1:]}" if ask else "") + (f" (ask: {who})" if who else "")
            doc.add_comment(runs, text, author=author, initials=author[:2])
            n_comments += 1

    out.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="tracked-") as tmp:
        built_path = Path(tmp) / out.name
        doc.save(built_path)
        with zipfile.ZipFile(template) as original, zipfile.ZipFile(built_path) as built, \
                zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as final:
            for entry in original.infolist():
                same = entry.filename not in EDITED
                final.writestr(entry, original.read(entry.filename) if same else built.read(entry.filename))
            for name in sorted(set(built.namelist()) - set(original.namelist())):
                final.writestr(built.getinfo(name), built.read(name))
    if template.read_bytes() != before_bytes:
        raise SystemExit("The template changed while building; rebuild from its new version.")
    return {"edits": len(edits), "comments": n_comments}
