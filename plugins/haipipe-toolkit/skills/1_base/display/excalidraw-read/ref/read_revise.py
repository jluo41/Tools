"""Read a file we revise (a co-author's Word file, through the PDF Word makes of it) section by section: one
picture per section cut from the PDF, the words each planned change touches under a see-through red box, the
updated draft's new words under a green one, and a note per change: the change, why, and in red the evidence
needed and who can answer, written smaller than the pages. The plan in one picture comes first (excalidraw-read's
revise mode; JL 261011: "take screen shot of the pdf version for each section and then say what to modify and
what to change").

    python read_revise.py draw <face.md> --key <plan key> [--out <drawing>.excalidraw] [--pdf <file>.pdf]

The plan is a yaml block in a studio topic's face (spec: ref/revise-spec.md): `file` (the .docx, beside the
face or in the Block's related/materials/), optional `pdf` (default: the .docx's name with .pdf, saved from
Word so the pages look as the author sees them), optional `stop_at` (the words that end the last section) and
`people` ({name: role}), and `edits`, each {id, part, title, after?, find, new?, image?, why, evidence?}. A section is
the title block or one Heading 1 of the .docx; `part` names it ("1".."9", "title", "R" for References, "F" for
Overview Figure, or the heading's text). The build stops when a find is not in its section of the PDF.
Needs poppler (pdftotext, pdftoppm), Pillow, PyYAML and python-docx; writes through haipipe-studio's canvas.write.
"""
import argparse
import re
import subprocess
import sys
import tempfile
from pathlib import Path

import yaml
from PIL import Image

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import read_paper as rp  # noqa: E402  (its drawing helpers, and haipipe-studio's scripts on sys.path)
import canvas  # noqa: E402

DPI = 200                      # the cut pictures
X0, PIC2, GAP2, NOTE2 = 80, 1000, 110, 360     # the section frame: the file today | updated draft | notes
LABEL = 24                                     # the change numbers in the margins
NCOL, CGAP = 3, 16                             # notes in three columns, so each stays level with its words
NOTE_TITLE, NOTE_BODY = 16, 14                 # notes smaller than the pages they explain (JL 261011: "The
QUOTE = 24                                     # content or the font size, can be much smaller, than the main
                                               # content"); a change's note quotes at most QUOTE of its words
X1 = X0 + PIC2 + GAP2
NX2 = X1 + PIC2 + 40
W2 = NX2 + NCOL * NOTE2 + (NCOL - 1) * CGAP + 40
GREEN = "#2f9e44"
OLD, NEW, SAME = (255, 196, 196), (190, 235, 198), (220, 220, 220)   # highlighter tints: removed, added, comment only
NAME = {"title": "Title and lead", "R": "References", "F": "Overview Figure"}
WORD = re.compile(r'xMin="([\d.]+)" yMin="([\d.]+)" xMax="([\d.]+)" yMax="([\d.]+)">([^<]*)</word>')


def tokens(s):
    return re.findall(r"[a-z0-9]+", s.lower())


def read_pages(pdf):
    """[(width, height, [(x0, y0, x1, y1, word), ...]), ...] in PDF points, one entry per page."""
    out = []
    for chunk in rp.run(["pdftotext", "-bbox", str(pdf), "-"]).split("<page ")[1:]:
        w, h = (float(re.search(rf'{k}="([\d.]+)"', chunk).group(1)) for k in ("width", "height"))
        out.append((w, h, [(float(a), float(b), float(c), float(d), t) for a, b, c, d, t in WORD.findall(chunk)]))
    return out


def edge_lines(h, ws):
    """{height: the whole line's text, page numbers as #} for the lines near the page's top or bottom edge."""
    lines = {}
    for _, y0, _, y1, t in ws:
        if y0 < 0.12 * h or y1 > 0.88 * h:
            lines.setdefault(round(y0), []).append(t)
    return {y: re.sub(r"\d+", "#", " ".join(ts)) for y, ts in lines.items()}


def page_words(pages):
    """Every body token in reading order: (page, x0, y0, x1, y1, token). The running header and footer, a whole
    line at the same height with the same text (page numbers aside) on two pages or more, are left out, so a
    section or a change that crosses a page never reads them."""
    import html
    from collections import Counter
    seen = Counter(k for _, h, ws in pages for k in edge_lines(h, ws).items())
    skip = {k for k, n in seen.items() if n >= 2}
    out = []
    for p, (_, h, ws) in enumerate(pages, 1):
        here = {y for y, text in edge_lines(h, ws).items() if (y, text) in skip}
        out += [(p, x0, y0, x1, y1, tok) for x0, y0, x1, y1, t in ws if round(y0) not in here
                for tok in tokens(html.unescape(t))]
    return out


def cut(pdf, pages, first, last, out):
    """Cut from `first` (page, y0) to `last` (page, y1), halfway between lines, stitched across pages. Returns the
    picture's size, a function (page, x, y) -> pixel (x, y) in it, and where it is ("page 2", "pages 2-3")."""
    k = DPI / 72
    parts, offsets = [], {}
    span = [w for p in range(first[0], last[0] + 1) for w in pages[p - 1][2]]
    left = max(0, min(w[0] for w in span) - 12) * k            # the sides: 12 pt outside the pages' own text, so a
    right = (max(w[2] for w in span) + 12) * k                 # narrow-margin file keeps its first letters
    with tempfile.TemporaryDirectory(prefix="revise-page-") as tmp:
        for p in range(first[0], last[0] + 1):
            stem = Path(tmp) / f"p{p}"
            rp.run(["pdftoppm", "-f", str(p), "-l", str(p), "-r", str(DPI), "-png", "-singlefile", str(pdf), str(stem)])
            im = Image.open(f"{stem}.png").convert("RGB")
            on = pages[p - 1][2]
            above = [w[3] for w in on if w[3] <= first[1]]
            below = [w[1] for w in on if w[1] >= last[1]]
            # a page's own text, without its running header and footer (the outer 32 pt): a narrow-margin file
            # starts its text above 0.6 inch, so a fixed trim would cut its first line
            body = [w for w in on if w[1] > 32 and w[3] < pages[p - 1][1] - 32]
            top = (int(min(w[1] for w in body) * k) - 10 if body else int(0.6 * DPI)) if p != first[0] else int(
                ((max(above) + first[1]) / 2 if above else first[1] - 5) * k)
            bottom = (int(max(w[3] for w in body) * k) + 10 if body else im.height - int(0.6 * DPI)) \
                if p != last[0] else int(((min(below) + last[1]) / 2 if below else last[1] + 5) * k)
            offsets[p] = (sum(x.height + 30 for x in parts), top)
            parts.append(im.crop((int(left), top, min(im.width, int(right)), bottom)))
    pic = Image.new("RGB", (max(x.width for x in parts), sum(x.height for x in parts) + 30 * (len(parts) - 1)), "white")
    for x, (y, _) in zip(parts, offsets.values()):
        pic.paste(x, (0, y))
    pic.save(out)
    return pic.size, (lambda p, x, y: (x * k - int(left), y * k - offsets[p][1] + offsets[p][0])), (
        f"page {first[0]}" if first[0] == last[0] else f"pages {first[0]}-{last[0]}")


def line_boxes(ws, to_px):
    """One box per printed line of a passage, in the picture's pixels: [(x0, y0, x1, y1)], top to bottom."""
    lines = {}
    for w in ws:
        lines.setdefault((w[0], next((t for p, t in lines if p == w[0] and abs(t - w[2]) < 3), w[2])), []).append(w)
    out = []
    for (p, _), ln in lines.items():
        x0, y0 = to_px(p, min(w[1] for w in ln), min(w[2] for w in ln))
        x1, y1 = to_px(p, max(w[3] for w in ln), max(w[4] for w in ln))
        out.append((x0 - 3, y0 - 2, x1 + 3, y1 + 2))
    return sorted(out, key=lambda r: (r[1], r[0]))


def highlight(png, marks):
    """Lay see-through boxes on a picture, [(boxes, tint)], as a marker pen does: the tint is multiplied onto
    the page, so the words stay black and readable under it (JL 261011: "just use the boxes with transparent
    to highlight the text part"); paper mode marks its quotes the same way."""
    from PIL import ImageChops, ImageDraw
    im = Image.open(png).convert("RGB")
    for boxes, tint in marks:
        mask = Image.new("RGB", im.size, (255, 255, 255))
        d = ImageDraw.Draw(mask)
        for box in boxes:
            d.rectangle(box, fill=tint)
        im = ImageChops.multiply(im, mask)
    im.save(png)


def split_change(old, new):
    """The smallest change on word edges: (kept before, removed, added, kept after)."""
    a = 0
    while a < min(len(old), len(new)) and old[a] == new[a]:
        a += 1
    b = 0
    while b < min(len(old), len(new)) - a and old[len(old) - 1 - b] == new[len(new) - 1 - b]:
        b += 1

    def inside(s, i):
        return 0 < i < len(s) and s[i - 1].isalnum() and s[i].isalnum()
    while a > 0 and (inside(old, a) or inside(new, a)):
        a -= 1
    while b > 0 and (inside(old, len(old) - b) or inside(new, len(new) - b)):
        b -= 1
    return old[:a], old[a:len(old) - b], new[a:len(new) - b], old[len(old) - b:]


def margin_numbers(labels, edge):
    """Draw the change numbers right-aligned against a picture's left edge, each centred on its first line, [(the
    line's middle, id, color)]: numbers on one printed line share one label ("10 11"); a number on the next line, too close to sit
    under the one above, steps one place to the left instead of moving off its line."""
    rows = []
    for y, n, color in sorted(labels, key=lambda r: r[0]):
        if rows and abs(y - rows[-1][0]) < 8:
            rows[-1][1].append(n)
        else:
            rows.append([y, [n], color])
    last, step = None, 0
    for y, ns, color in rows:
        step = 1 - step if last is not None and y - last < LABEL * 1.25 else 0
        s, last = " ".join(ns), y
        rp.text(edge - 8 - len(s) * LABEL * 0.55 - step * 40, y - LABEL * 0.625, s, LABEL, color)


def quoted(s):
    w = s.split()
    return " ".join(w[:QUOTE]) + (" ..." if len(w) > QUOTE else "")


def note(x, y, s, width, size, color=rp.INK):
    """A side note's lines, smaller than the page pictures beside it (canvas.aside); returns the y below it."""
    t = canvas.aside(rp.text(x, y, rp.wrap(s, width, size), size, color))
    return y + t["height"]


def change_text(e):
    new = e.get("new")
    if e.get("image"):
        return "Put our figure in the box; the box's instructions give way to it."
    if new is None:
        return "No change to the words; a comment asks for the evidence."
    keep_a, removed, added, _ = split_change(e["find"], new)
    removed, added = (quoted(" / ".join(x.split("\n"))) for x in (removed, added))   # a new line shows as " / "
    if removed.strip() and added.strip():
        return f"\"{removed.strip()}\" becomes \"{added.strip()}\""
    if removed.strip():
        return f"Delete \"{removed.strip()}\"."
    return f"After \"...{' '.join((keep_a or e['find']).split()[-5:])}\" add: \"{added.strip()}\""


def load(face, key):
    face = Path(face).resolve()
    for block in re.findall(r"```yaml\n(.*?)\n```", face.read_text(encoding="utf-8"), re.S):
        if re.search(rf"(?m)^{key}:", block):
            plan = yaml.safe_load(block)[key]
            break
    else:
        raise SystemExit(f"{face.name} has no yaml block {key}")
    board = next((p for p in face.parents if (p / "board.md").is_file()), face.parent)
    docx = next((p for p in (face.parent / plan["file"], board / "related/materials" / plan["file"]) if p.is_file()),
                None)
    if docx is None:
        raise SystemExit(f"{plan['file']} is neither beside the face nor in {board.name}/related/materials/")
    return face, plan, docx


def sections_of(docx):
    """[(id, heading text)]: the title block, then each Heading 1, in the file's order."""
    from docx import Document
    doc = Document(docx)
    title = next((p.text.strip() for p in doc.paragraphs if re.match(r"Task \d+(\.\d+)?:", p.text.strip())), None)
    out = [("title", title)] if title else []
    for p in doc.paragraphs:
        if p.style.name == "Heading 1" and p.text.strip():
            t = re.sub(r"\s+", " ", p.text.strip())
            m = re.match(r"(\d+)\s", t)
            out.append((m.group(1) if m else {"References": "R", "Overview Figure": "F"}.get(t, t), t))
    return out


ABBREV = {"al", "e.g", "i.e", "vs", "dr", "fig", "no", "approx", "etc"}   # a full stop after these ends no sentence
WNS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def split_sentences(text):
    """[(start, end)] of each sentence in `text`: a break after . ? or ! before a capital or an opening
    bracket, but not after an abbreviation ("et al.", "e.g.")."""
    out, start = [], 0
    for m in re.finditer(r"(?<=[.!?])[\"”)]?\s+(?=[A-Z(\[“\"])", text):
        word = re.findall(r"([A-Za-z.]+)\.\S*$", text[start:m.start()].rstrip())
        if word and word[-1].lower().rstrip(".") in ABBREV:
            continue
        out.append((start, m.start()))
        start = m.end()
    if text[start:].strip():
        out.append((start, len(text)))
    return [(a, b) for a, b in out if text[a:b].strip()]


def draft_units(draft, stop_at):
    """{section id: [(sentence, changed)]}: the tracked Word file read with every change accepted, its
    guidance (italics) left out, each paragraph cut into sentences, each table row one unit ("When ·
    Milestone"), each reference one unit. `changed` when a tracked change falls inside the unit."""
    import zipfile
    from lxml import etree
    root = etree.fromstring(zipfile.ZipFile(draft).read("word/document.xml"))
    W = WNS

    def accepted(p):   # the paragraph's text with changes accepted; inserted characters; where words were deleted
        text, ins, dels = "", set(), set()
        whole = p.find(f"{W}pPr/{W}rPr/{W}ins") is not None
        for e in p.iter(W + "t", W + "del", W + "br"):
            if e.tag == W + "del":
                dels.add(len(text))
            elif e.tag == W + "br":                     # a line break inside a cell or paragraph is a space,
                if not any(a.tag == W + "del" for a in e.iterancestors()):   # not two words run together
                    text += " "
            elif not any(a.tag == W + "del" for a in e.iterancestors()):
                if whole or any(a.tag == W + "ins" for a in e.iterancestors()):
                    ins.update(k for k in range(len(text), len(text) + len(e.text or "")) if not e.text[k - len(text)].isspace())
                text += e.text or ""
        return text, ins, dels

    def guidance(p):
        runs = [r for r in p.iter(W + "r") if "".join(t.text or "" for t in r.iter(W + "t")).strip()]
        return bool(runs) and all(r.find(f"{W}rPr/{W}i") is not None for r in runs)

    out, sid, para = {}, None, {}
    body = root.find(W + "body")
    for el in body:
        if el.tag == W + "tbl":
            if sid is None:
                continue
            for row in el.findall(W + "tr")[1:]:            # the first row is the table's header
                cells = [accepted(p) for p in row.iter(W + "p") if not guidance(p)]
                text = " · ".join(t for t, _, _ in cells if t.strip())
                if text.strip():
                    out[sid].append((text, any(i or d for _, i, d in cells)))
                    para.setdefault(sid, {})[len(out[sid]) - 1] = ("table", id(el))
            continue
        if el.tag != W + "p":
            continue
        text, ins, dels = accepted(el)
        clean = re.sub(r"\s+", " ", text).strip()
        if stop_at and stop_at in clean:
            break
        st = el.find(f"{W}pPr/{W}pStyle")
        if st is not None and st.get(W + "val") in ("Heading1", "Heading 1") and clean:
            m = re.match(r"(\d+)\s", clean)
            sid = m.group(1) if m else {"References": "R", "Overview Figure": "F"}.get(clean, clean)
            out[sid] = []
            continue
        if re.match(r"Task \d+(\.\d+)?:", clean):
            sid = "title"
            out[sid] = []
        if sid is None or not clean or guidance(el):
            continue
        spans = [(0, len(text))] if sid == "R" else split_sentences(text)
        while len(spans) > 1 and len(text[spans[0][0]:spans[0][1]].split()) == 1:   # a label ("Figure.") joins
            spans[:2] = [(spans[0][0], spans[1][1])]                                # the sentence after it
        for a, b in spans:
            changed = any(a <= k < b for k in ins) or any(a < k <= b for k in dels)
            out[sid].append((re.sub(r"\s+", " ", text[a:b]).strip(), changed))
            para.setdefault(sid, {})[len(out[sid]) - 1] = ("p", id(el))
    DRAFT_PARAS.clear()
    DRAFT_PARAS.update(para)
    return out


DRAFT_PARAS = {}                        # draft_units: {section id: {unit index: its paragraph}}
PARA_BRIEF = {}                         # read_draft: {"C3.P5": the paragraph's one-line brief}
PAGE = HERE.parents[2] / "page" / "haipipe-draft"     # haipipe-draft: its three-section draft grammar


def plan_layout():
    sys.path.insert(0, str(PAGE))
    from src import plan_layout as pl
    return pl


def draft_file(path):
    """The current version in a Page-style draft/ folder (haipipe-draft): the one `<stem>-draft-v<G>.<S>.md`
    in it; older versions sit in previous/, records in records/."""
    path = Path(path)
    if path.is_file():
        return path
    found = sorted(path.glob("*-draft-v*.md"))
    if len(found) != 1:
        raise SystemExit(f"{path}: {len(found)} current drafts; keep one, move older ones to previous/")
    return found[0]


def section_ids(docx):
    """{a division's name, lower case: the .docx section id}, so `## C2 · Challenge and Significance` is section 1."""
    return {re.sub(r"^\d+\s+", "", NAME.get(sid, head)).lower(): sid for sid, head in sections_of(docx)}


def read_draft(path, docx):
    """{section id: [(S number, [change ids], sentence, address)]} from the draft: haipipe-draft's three-section
    Markdown, read through its own `to_canonical`. A Bullet `- B<k> · S<n> · <sentence>` is one sentence; its
    `Change:` line names the planned changes that touch it."""
    canon = plan_layout().to_canonical(draft_file(path).read_text(encoding="utf-8"))
    names, out, sid, para, cur = section_ids(docx), {}, None, None, None
    for line in canon.split("\n"):
        m = re.match(r"^## (C\d+) · (.+)$", line)
        if m:
            sid, cur = names.get(m.group(2).strip().lower()), None
            if sid is not None:
                out[sid] = []
            continue
        m = re.match(r"^### (C\d+\.P\d+)(?: · (.*\S))?", line)
        if m:
            para, cur = m.group(1), None
            PARA_BRIEF[para] = m.group(2) or ""
            continue
        m = re.match(r"^- B(\d+) · S(\d+) · (.*\S)\s*$", line)
        if m and sid is not None:
            cur = [int(m.group(2)), [], m.group(3), f"{para}.B{m.group(1)}"]
            out[sid].append(cur)
            continue
        m = re.match(r"^\s+Change: (.*)$", line)
        if m and cur is not None:
            cur[1] = re.findall(r"(?:^|,\s*)(\w+) · ", m.group(1)) or [x.strip() for x in m.group(1).split(",")]
    return {k: [tuple(u) for u in v] for k, v in out.items()}


def cmd_seed(a):
    """Write the topic's draft once (draft/<stem>-draft-v1.0.md, haipipe-draft's three sections) from the tracked
    Word file (plan `draft`) with every change accepted: one Bullet per sentence, a paragraph per paragraph of
    the file, each sentence's `Change:` line naming the changes whose added words fall in it. A person or an
    agent then writes each Bullet's Point. Never overwrites a draft that exists."""
    import datetime
    face, plan, docx = load(a.face, a.key)
    board = next((p for p in face.parents if (p / "board.md").is_file()), face.parent)
    if not plan.get("sentences") or not plan.get("draft"):
        raise SystemExit("the plan needs `sentences` (the draft folder beside the face) and `draft` (the tracked file)")
    folder = face.parent / plan["sentences"]
    if list(folder.glob("*-draft-v*.md")):
        raise SystemExit(f"{folder} already holds a draft; it is the source now")
    units = draft_units(board / plan["draft"], plan.get("stop_at"))
    titles = {str(e["id"]): e["title"] for e in plan["edits"]}
    stem = plan.get("stem") or re.sub(r"[^a-z0-9]", "", Path(plan["file"]).stem.split(" - ")[0].lower())
    out = [f"# {stem} · draft v1.0", "draft-version: v1.0", "supersedes: none",
           f"date: {datetime.date.today():%y%m%d}", "approved: ⬜",
           f"arc: {plan.get('arc', 'The updated text of ' + plan['file'] + ', one sentence per Bullet.')}", ""]
    c = p = 0
    for sid, head in sections_of(docx):
        sents = units.get(sid) or []
        if not sents:
            continue
        spans, toks = [], []                            # each sentence's token range in the section's stream
        for sent, _ in sents:
            t = tokens(sent)
            spans.append((len(toks), len(toks) + len(t)))
            toks += t
        ids = [[] for _ in sents]
        for e in plan["edits"]:
            if str(e["part"]) not in (sid, head) or e.get("new") is None or e.get("image"):
                continue                                # a part may be named by id ("R") or heading ("References")
            added, whole = tokens(split_change(e["find"], e["new"])[2]), tokens(e["new"])
            # the added words inside the change's whole new text, so a one-word change ("Q4") lands in its row
            w0 = next((m for m in range(len(toks) - len(whole) + 1) if whole and toks[m:m + len(whole)] == whole), None)
            off = next((m for m in range(len(whole) - len(added) + 1) if added and whole[m:m + len(added)] == added), 0)
            at = None if w0 is None or not added else w0 + off
            if at is None and w0 is not None and whole:  # a change of punctuation only (",") adds no word:
                at, added = w0, whole                     # tag the sentence its new text sits in
            for k, (lo, hi) in enumerate(spans):
                if at is not None and lo < at + len(added) and at < hi:
                    ids[k].append(str(e["id"]))
        c += 1
        name = re.sub(r"^\d+\s+", "", NAME.get(sid, head))
        out += [f"## C{c} · {name}", ""]
        last, b = None, 0
        for k, (sent, _) in enumerate(sents):
            here = DRAFT_PARAS.get(sid, {}).get(k)
            if here != last:
                p, b, last = p + 1, 0, here
                out += ([""] if k else []) + [f"### C{c}.P{p} · " + " ".join(sent.split()[:6]) + " ..."]
            b += 1
            out += [f"- B{b} · S{k + 1} · {sent}",
                    "  Point: " + (f"[Change] {titles[ids[k][0]]}" if ids[k] else "[Original] The file author's sentence")]
            if ids[k]:
                out.append("  Change: " + ", ".join(f"{i} · {titles[i]}" for i in ids[k]))
        out.append("")
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{stem}-draft-v1.0.md"
    path.write_text(plan_layout().to_sectioned("\n".join(out) + "\n"), encoding="utf-8")
    print(f"wrote {path.name}: {sum(len(v) for v in units.values())} sentences; write each Bullet's Point next")


def bounds_of(words, docx, stop_at, pdf):
    """[(id, heading, first word, end word)] for each section of `docx` as `pdf` prints it."""
    toks = [w[5] for w in words]

    def first_after(phrase, start):
        want = tokens(phrase)
        for i in range(start, len(toks) - len(want) + 1):
            if toks[i:i + len(want)] == want:
                return i
        raise SystemExit(f"not found in {pdf.name} after word {start}: {phrase!r}")
    marks, k = [], 0
    for sid, head in sections_of(docx):
        k = first_after(head, k)
        marks.append((sid, head, k))
    end = first_after(stop_at, k) if stop_at else len(words)
    return [(sid, head, i, marks[n + 1][2] if n + 1 < len(marks) else end) for n, (sid, head, i) in enumerate(marks)]


def find_in(sec, phrase, after, what):
    """The words of `phrase` in a section: the first match after each `after` phrase in turn."""
    stoks, start = [w[5] for w in sec], 0
    for ph in after or []:
        aw = tokens(str(ph))
        at = next((m for m in range(start, len(stoks) - len(aw) + 1) if stoks[m:m + len(aw)] == aw), None)
        if at is None:
            raise SystemExit(f"{what}: after-phrase {ph!r} not in its section")
        start = at + len(aw)
    want = tokens(phrase)
    hit = next((m for m in range(start, len(stoks) - len(want) + 1) if stoks[m:m + len(want)] == want), None)
    if hit is None and after:
        # a table row's label printed among its cell's lines (LibreOffice prints a two-line cell's "Year 3"
        # between the lines; b02 261011): look again with the label's words, those not in the phrase, left out
        drop = {t for ph in after for t in tokens(str(ph))} - set(want)
        idx = [m for m in range(start, len(stoks)) if stoks[m] not in drop]
        sub = [stoks[m] for m in idx]
        k = next((m for m in range(len(sub) - len(want) + 1) if sub[m:m + len(want)] == want), None)
        if k is not None:
            return idx[k], [sec[m] for m in idx[k:k + len(want)]]
    if hit is None:
        # a few other words inside the phrase: a row split across pages, with the page's header and footer and
        # the table's repeated header row between its lines (LibreOffice; b02 261011). The phrase's words in
        # order, at most GAP other words between two of them
        GAP = 14
        for m in range(start, len(stoks)):
            if stoks[m] != want[0]:
                continue
            got, at = [m], m
            for t in want[1:]:
                nxt = next((q for q in range(at + 1, min(at + 2 + GAP, len(stoks))) if stoks[q] == t), None)
                if nxt is None:
                    break
                got.append(nxt)
                at = nxt
            if len(got) == len(want):
                return m, [sec[q] for q in got]
    if hit is None:
        raise SystemExit(f"{what}: {phrase!r} not in its section")
    return hit, sec[hit:hit + len(want)]


def evidence_of(e):
    evs = e.get("evidence") or []
    return evs if isinstance(evs, list) else [evs]


def draw_plan(plan, face, docx, draft, edits, counts, people):
    """One screen: where the file goes, words per part before and after, and who answers the comments."""
    import datetime
    fr = rp.frame(plan.get("plan_title", f"{docx.stem}: the plan in one picture"), 2400, 1300)
    y = rp.para(40, 30, plan.get("plan_heading", f"{docx.stem}: the plan in one picture"), 2320, 34) + 10
    y = rp.para(40, y, "The file keeps its shape, fonts, cover and comments; every change is a Word tracked change "
                       f"by {plan.get('author', 'JL')} that the file's owner accepts or rejects, and a comment sits "
                       "only on the few words someone must answer.", 2320, 20) + 40
    saved = datetime.datetime.fromtimestamp(docx.stat().st_mtime).strftime("%Y-%m-%d %H:%M")
    n_ev = sum(len(evidence_of(e)) for e in edits)
    boxes = [("The shared file", plan.get("source", "the original, read only"), rp.INK),
             ("The copy we read", f"related/materials/{docx.name}, saved {saved}", rp.INK),
             ("The plan of changes", f"{len(edits)} changes in {face.parent.name}", rp.INK),
             ("The tracked Word file", f"{draft.name if draft else 'not built yet'}: {len(edits)} changes, "
                                       f"{n_ev} comments", rp.INK),
             ("? Back to the shared folder", plan.get("send", "a person copies it beside the original"), rp.RED)]
    bw, gap = 400, 76
    for n, (title, body, color) in enumerate(boxes):
        x = 40 + n * (bw + gap)
        yy = rp.para(x + 16, y + 14, title, bw - 32, 24, color) + 8
        yy = rp.para(x + 16, yy, body, bw - 32, 20, color)
        rp.base("rectangle", x, y, bw, 200, color, rough=0)
        if n:
            rp.arrow(x - gap + 6, y + 100, x - 6, y + 100)
    y += 260
    top = y
    y = rp.para(40, y, "Words per part as printed: the file today (thin) and the updated draft (bold)", 1100, 26) + 20
    most = max(max(a, b) for _, a, b in counts) or 1
    for name, a, b in counts:
        rp.para(40, y, name, 340, 20)
        for k, (v, sw) in enumerate(((a, 1), (b, 3))):
            w = max(4, 700 * v / most)
            rp.base("rectangle", 400, y + k * 24, w, 18, rp.INK, sw=sw, rough=0)
            rp.text(410 + w, y + k * 24 - 4, str(v), 20)
        y += 64
    y2 = rp.para(1260, top, "Who answers the comments", 1100, 26, rp.RED) + 14
    names = list(people) + sorted({w for e in edits for ev in evidence_of(e) for w in ev.get("who", [])} - set(people))
    for person in names:
        mine = [f"{e['id']}: {ev.get('ask') or ev['need']}" for e in edits for ev in evidence_of(e)
                if person in ev.get("who", [])]
        if mine:
            y2 = rp.para(1260, y2, person + (f", {people[person]}" if person in people else ""), 1100, 22, rp.RED) + 4
            for line in mine:
                y2 = rp.para(1290, y2, line, 1070, 20, rp.RED) + 2
            y2 += 10
    for q in plan.get("open") or []:
        y2 = rp.para(1260, y2 + 6, "? " + q, 1100, 20, rp.RED)
    fr["height"] = max(y, y2) + 40
    return fr


def cmd_draw(a):
    face, plan, docx = load(a.face, a.key)
    pdf = Path(a.pdf) if a.pdf else docx.parent / (plan.get("pdf") or docx.with_suffix(".pdf").name)
    board = next((p for p in face.parents if (p / "board.md").is_file()), face.parent)
    draft = board / plan["draft"] if plan.get("draft") else None
    # sentence mode: the middle column is the topic's sentence draft, as text a person edits (JL 261011: "we
    # don't have this, but use the text to replace it"; "The Sentences of para, should come from the draft of
    # the sNN"); otherwise it is the updated draft's printed pages
    sfile = draft_file(face.parent / plan["sentences"]) if plan.get("sentences") else None
    sents = read_draft(sfile, docx) if sfile else {}
    upd = None if sfile else (Path(a.updated_pdf) if a.updated_pdf else (draft.with_suffix(".pdf") if draft else None))
    for f, what in ((pdf, docx.name),) + (() if sfile else ((upd, draft.name if draft else "the updated draft"),)):
        if f is None or not f.is_file():
            raise SystemExit(f"Save {what} as PDF from Word first (File > Save As > PDF, changes accepted for the "
                             f"draft), as {f}")
    pages = read_pages(pdf)
    upages = read_pages(upd) if upd else []
    words = page_words(pages)
    uwords = page_words(upages) if upd else []
    secs = bounds_of(words, docx, plan.get("stop_at"), pdf)
    usecs = {sid: (i, j) for sid, _, i, j in bounds_of(uwords, docx, plan.get("stop_at"), upd)} if upd else {}
    edits = plan["edits"]
    people = plan.get("people") or {}
    after = (lambda sid: sum(len(u[2].split()) for u in sents.get(sid, []))) if sfile else (
        lambda sid: usecs[sid][1] - usecs[sid][0])
    counts = [(NAME.get(sid, head), j - i, after(sid)) for sid, head, i, j in secs if sid not in ("title",)]
    plan_frame = draw_plan(plan, face, docx, draft if draft and draft.is_file() else None, edits, counts, people)
    side = rp.frame(plan.get("title", "The file today and the updated draft, section by section"), W2)
    canvas.reads_down(side)
    y = rp.para(40, 24, plan.get("heading", f"{docx.stem} and the updated draft, section by section"), W2 - 80, 34) + 14
    middle = ("Middle: the updated draft, one sentence per box; green, with the change's number, where a planned "
              "change touches it; double-click a sentence to change it." if sfile else
              "Middle: the updated draft with every tracked change accepted; green highlight: the new words.")
    y = rp.para(40, y, f"Left: the file today; red highlight: the words that change. {middle} Each column names its "
                       "file under its heading. Right: why, what was found, and in red the evidence needed and who can "
                       "answer; a note in the first column has an arrow, one further right sits at the height of its "
                       "words and carries the same number." + (f" {plan['intro']}" if plan.get("intro") else ""),
                W2 - 80, 20) + 24
    rp.text(X0, y, plan.get("left_label", "The file today"), 28)
    rp.text(X1, y, "Updated draft, sentence by sentence" if sfile else "Updated draft, changes accepted", 28, GREEN)
    rp.text(NX2, y, "Change · why · who answers", 28)
    # which file each picture column shows, relative to the Block (JL 261011: "you should put which files we are
    # using here, for the two columns")
    rel = lambda f: str(Path(f).resolve().relative_to(board)) if Path(f).resolve().is_relative_to(board) else Path(f).name
    rp.text(X0, y + 40, f"{rel(docx)}\n→ printed as {pdf.name}", 20, rp.INK, rp.CODE)
    if sfile:
        rp.text(X1, y + 40, f"{rel(sfile)}\n→ the draft (haipipe-draft); the Word file is built from it last", 20,
                GREEN, rp.CODE)
    else:
        rp.text(X1, y + 40, f"{plan['draft'] if plan.get('draft') else upd.name}\n→ every tracked change accepted, "
                            "then printed", 20, GREEN, rp.CODE)
    y += 52
    marks = plan.get("change_notes") or []            # green ✎ notes: at a column (left, middle, notes) or a change
    heads = {"left": X0, "middle": X1, "notes": NX2}
    for c in marks:
        if str(c["at"]) in heads:
            rp.els.append(canvas.change_note(heads[str(c["at"])], y + 40, c["text"], str(c.get("date") or "") or None,
                                             rp.FRAME[0], 16))
    y += 56 + (30 if any(str(c["at"]) in heads for c in marks) else 0)
    done = set()
    units = draft_units(draft, plan.get("stop_at")) if upd and draft and draft.is_file() else {}
    with tempfile.TemporaryDirectory(prefix="revise-") as tmp:
        for n, (sid, head, i, j) in enumerate(secs):
            mine = [e for e in edits if str(e["part"]) in (sid, head)]
            if not mine:
                continue
            sec = words[i:j]
            usec = uwords[usecs[sid][0]:usecs[sid][1]] if upd else []
            shots = []
            for tag, src, pg, ws in (("l", pdf, pages, sec),) + ((("r", upd, upages, usec),) if upd else ()):
                out = Path(tmp) / f"{tag}-{n}.png"
                shots.append(cut(src, pg, (ws[0][0], ws[0][2]), (ws[-1][0], ws[-1][4]), out) + (out,))
            (lw, lh), lpx, lwhere, lout = shots[0]
            if upd:
                (uw, uh), upx, uwhere, uout = shots[1]
            rows = []
            for e in mine:
                hit, ws = find_in(sec, e["find"], e.get("after"), f"change {e['id']} in {pdf.name}")
                new, uws, tint = e.get("new"), [], NEW
                if not e.get("image") and upd:
                    if new is None:                      # a comment only: the same words, in gray
                        uws, tint = find_in(usec, e["find"], e.get("after"), f"change {e['id']}")[1], SAME
                    else:
                        _, removed, added, _ = split_change(e["find"], new)
                        if added.strip():
                            k, span = find_in(usec, new, e.get("after"), f"change {e['id']} in {upd.name}")
                            inner = tokens(added)
                            st = [w[5] for w in span]
                            at = next((m for m in range(len(st) - len(inner) + 1) if st[m:m + len(inner)] == inner), 0)
                            uws = span[at:at + len(inner)]
                rows.append((hit, e, line_boxes(ws, lpx), line_boxes(uws, upx) if uws else [], tint))
            highlight(lout, [(lb, OLD) for _, _, lb, _, _ in rows])
            rp.text(X0, y, f"{NAME.get(sid, head)} · today, {lwhere}", 22)
            py = y + 36
            lbot = rp.picture(X0, py, lout, PIC2)
            ls = PIC2 / lw
            where = {}                                   # change id -> (top, bottom) of its first sentence
            if upd:
                highlight(uout, [(ub, t) for _, _, _, ub, t in rows if ub])
                rp.text(X1, y, f"updated, {uwhere}", 22, GREEN)
                ubot = rp.picture(X1, py, uout, PIC2)
                us = PIC2 / uw
            else:
                paras = len({u[3].rsplit(".", 1)[0] for u in sents.get(sid, [])})
                rp.text(X1, y, f"{NAME.get(sid, head)} · updated, {paras} paragraph{'s' * (paras != 1)}, "
                               f"{len(sents.get(sid, []))} sentences", 22, GREEN)
                ubot, where = sentence_column(sents.get(sid, []), sid, py)
            bottoms = [py] * NCOL                        # each note column's lowest point so far
            lnums, unums = [], []
            for hit, e, lb, ub, tint in sorted(rows, key=lambda r: r[0]):
                ltop, lbtm = py + lb[0][1] * ls, py + max(r[3] for r in lb) * ls
                lnums.append((py + (lb[0][1] + lb[0][3]) / 2 * ls, str(e["id"]), rp.RED))
                mid = (ltop + lbtm) / 2
                if ub:
                    utop, ubtm = py + ub[0][1] * us, py + max(r[3] for r in ub) * us
                    unums.append((py + (ub[0][1] + ub[0][3]) / 2 * us, str(e["id"]), rp.INK if tint == SAME else GREEN))
                    mid = (utop + ubtm) / 2
                elif str(e["id"]) in where:              # its sentence in the middle column
                    utop, ubtm = where[str(e["id"])]
                    unums.append((utop + 13, str(e["id"]), GREEN))
                    mid = (utop + ubtm) / 2
                # the column where the note sits highest, nearest its words; a tie goes to the left
                col = min(range(NCOL), key=lambda c: (max(bottoms[c], mid - 24), c))
                ny, nx = max(bottoms[col], mid - 24), NX2 + col * (NOTE2 + CGAP)
                yy = note(nx + 12, ny + 10, f"{e['id']} · {e['title']}", NOTE2 - 24, NOTE_TITLE) + 4
                if upd:                                  # in sentence mode the middle column shows the change
                    yy = note(nx + 12, yy, "Change: " + change_text(e), NOTE2 - 24, NOTE_BODY) + 4
                yy = note(nx + 12, yy, "Why: " + e["why"], NOTE2 - 24, NOTE_BODY) + 4
                for ev in evidence_of(e):
                    ask = ev.get("ask", "")
                    if ev.get("found"):                  # papers a Discovery Run found (haipipe-discovery)
                        yy = note(nx + 12, yy, "Found: " + "; ".join(f"{f['cite']} ({f['run']})" if f.get("run") else
                                                                     f["cite"] for f in ev["found"]),
                                  NOTE2 - 24, NOTE_BODY) + 4
                    yy = note(nx + 12, yy, f"Evidence needed: {ev['need']}. {ask[:1].upper() + ask[1:]} "
                                           f"({', '.join(ev.get('who', []))})", NOTE2 - 24, NOTE_BODY, rp.RED) + 4
                for c in marks:
                    if str(c["at"]) == str(e["id"]):
                        rp.els.append(canvas.change_note(nx + 12, yy, c["text"], str(c.get("date") or "") or None,
                                                         rp.FRAME[0], 14))
                        yy += 20
                rp.base("rectangle", nx, ny, NOTE2, yy - ny + 6, rp.INK, rough=0)
                if col == 0:                             # further right, an arrow would cross the first column
                    rp.arrow(X1 + PIC2 + 6, mid, nx - 6, ny + 16, rp.RED)
                bottoms[col] = yy + 18
                done.add(e["id"])
            margin_numbers(lnums, X0)
            margin_numbers(unums, X1 - 36 if not upd else X1)
            y = max(lbot, ubot) + 50
            if upd and units.get(sid):                   # the updated section, one editable box per sentence
                y = sentence_band(units[sid], sid, y)
            y = max(y, *bottoms) + 70
    missing = [e["id"] for e in edits if e["id"] not in done]
    if missing:
        raise SystemExit(f"changes whose part names no section of {docx.name}: {missing}")
    side["height"] = y + 30
    rp.FRAME[0] = None
    canvas.grid(rp.els, [[plan_frame["id"], side["id"]]])
    out = Path(a.out) if a.out else face.parent / f"{face.parent.name}.excalidraw"
    canvas.write(out, rp.els, "excalidraw-read/ref/read_revise.py", files=rp.FILES)
    png = out.with_suffix(".png")
    subprocess.run([sys.executable, str(rp.STUDIO / "render_png.py"), str(out), str(png)], check=False)
    print(f"drew the plan and {len(done)} changes over {len({str(e['part']) for e in edits})} sections: {pdf.name} "
          f"({len(pages)} pages) beside " + (f"{upd.name} ({len(upages)} pages)" if upd else
          f"{sfile.name} ({sum(len(v) for v in sents.values())} sentences)") + f" -> {out}")


BAND = X1 + PIC2 - X0                  # the sentence band spans both pictures


def sentence_band(sents, sid, y):
    """Under a section's pictures: its sentences as the updated draft has them, one text box each, so a person
    edits a sentence in place (JL 261011: "in this area, you will put each sentence for the above section";
    "sentence level editable text box, so I can change them directly"). Green: a tracked change falls in it.
    Each box keeps what was built in customData, so `read_revise.py edits` can list what a person changed.
    Returns the y below the band."""
    rp.text(X0, y, "Sentences of the updated draft · double-click a sentence to change it", 22)
    yy = top = y + 40
    for k, (sent, changed) in enumerate(sents, 1):
        color = GREEN if changed else rp.INK
        lab = f"S{k}"
        rp.text(X0 - 10 - len(lab) * 16 * 0.55, yy + 3, lab, 16, color)
        e = rp.text(X0 + 12, yy, rp.wrap(sent, BAND - 24, 20), 20, color)
        e.update(originalText=sent, autoResize=False, width=BAND - 24,
                 customData={"sentence": {"part": sid, "n": k, "built": sent, "changed": changed}})
        yy += e["height"] + 8
    rp.base("rectangle", X0, top - 10, BAND, yy - top + 14, rp.INK, rough=0)
    return yy + 10


def sentence_column(sents, sid, y):
    """The middle column in sentence mode: one box per paragraph of the draft, its label ("P5 · <brief>") above
    it, and in it one text box per sentence, so a person edits a sentence in place; green where a planned change
    touches it (JL 261011: "IT IS NOT GOOD WE HAVE ONE PARAGRAPH WITH 17 SENTENCES"). Each sentence keeps what
    was built in customData, so `read_revise.py edits` lists what a person changed. Returns (the y below the
    column, {change id: (top, bottom) of its first sentence})."""
    yy, where, last, top = y + 6, {}, None, None

    def close(top, bottom):
        rp.base("rectangle", X1, top - 8, PIC2, bottom - top + 14, rp.INK, rough=0)
    for n, ids, sent, address in sents:
        para = address.rsplit(".", 1)[0]
        if para != last:
            if last is not None:
                close(top, yy)
                yy += 24
            label = para.split(".")[1] + (f" · {PARA_BRIEF[para]}" if PARA_BRIEF.get(para) else "")
            rp.text(X1 + 4, yy, label, 18, rp.INK)
            yy, top, last = yy + 32, yy + 32, para
        color = GREEN if ids else rp.INK
        lab = f"S{n}"
        rp.text(X1 - 8 - len(lab) * 16 * 0.55, yy + 3, lab, 16, color)
        e = rp.text(X1 + 12, yy, rp.wrap(sent, PIC2 - 24, 20), 20, color)
        e.update(originalText=sent, autoResize=False, width=PIC2 - 24,
                 customData={"sentence": {"part": sid, "n": n, "built": sent, "changed": bool(ids), "ids": ids,
                                          "address": address}})
        for i in ids:
            where.setdefault(i, (yy, yy + e["height"]))
        yy += e["height"] + 8
    if last is not None:
        close(top, yy)
    return yy + 10, where


def sentence_text(e):
    return " ".join((e.get("originalText") or e.get("text") or "").split())


def cmd_edits(a):
    """List the sentences a person changed on the drawing: Before (as built) and After (as edited)."""
    import json
    face, plan, docx = load(a.face, a.key)
    out = Path(a.drawing) if a.drawing else face.parent / f"{face.parent.name}.excalidraw"
    els = json.loads(out.read_text())["elements"]
    rows = []
    for e in els:
        info = (e.get("customData") or {}).get("sentence")
        if e.get("type") == "text" and info and not e.get("isDeleted") and sentence_text(e) != info["built"]:
            rows.append((info["part"], info["n"], info["built"], sentence_text(e)))
    order = {sid: i for i, (sid, _) in enumerate(sections_of(docx))}
    rows.sort(key=lambda r: (order.get(r[0], 99), r[1]))
    if not rows:
        print(f"no sentence changed on {out.name}")
    if rows and getattr(a, "apply", False):
        # fold the edits into the sentence draft, then take the edited boxes off the drawing: the next build
        # draws them from the draft (else a rebuild keeps the person's box beside the new one)
        if not plan.get("sentences"):
            raise SystemExit("--apply needs the plan's `sentences` draft folder")
        sfile = draft_file(face.parent / plan["sentences"])
        lines = sfile.read_text(encoding="utf-8").split("\n")
        built = {(e["customData"]["sentence"].get("address"), e["customData"]["sentence"]["built"]): sentence_text(e)
                 for e in els if (e.get("customData") or {}).get("sentence") and not e.get("isDeleted")
                 and sentence_text(e) != e["customData"]["sentence"]["built"]}
        section, para, hit = None, None, 0
        for k, line in enumerate(lines):
            m = re.match(r"^## (\d+) · (\w+)", line)
            if m:
                section = m.group(2)
            m = re.match(r"^### (C\d+\.P\d+)", line)
            if m:
                para = m.group(1)
            m = re.match(r"^(- B(\d+) · S\d+ · )(.*\S)\s*$", line)
            if section == "Draft" and m and (f"{para}.B{m.group(2)}", m.group(3)) in built:
                lines[k] = m.group(1) + built[(f"{para}.B{m.group(2)}", m.group(3))]
                hit += 1
        if hit != len(built):
            raise SystemExit(f"{sfile.name}: {len(built) - hit} edited sentence(s) no longer read as built; apply by hand")
        sfile.write_text("\n".join(lines), encoding="utf-8")
        doc = json.loads(out.read_text())
        gone = {id(e) for e in els if (e.get("customData") or {}).get("sentence") and sentence_text(e) !=
                e["customData"]["sentence"]["built"]}
        doc["elements"] = [e for e in els if id(e) not in gone]
        out.write_text(json.dumps(doc, ensure_ascii=False, indent=1))
        print(f"applied {len(rows)} sentence edit(s) to {sfile.name}; rebuild the drawing")
    for part, n, before, after in rows:
        _, removed, added, _ = split_change(before, after)
        print(f"{NAME.get(part, part)} · S{n}\n  before: {before}\n  after:  {after}\n"
              f"  change: \"{removed.strip()}\" -> \"{added.strip()}\"\n")
    return rows


def office_pdf(docx, pdf):
    """Print a Word file to PDF with LibreOffice through the Codex document runtime's renderer, close to Word's own
    pages (b02 and b04 print this way; JL: "I don't want a new format"). The Word file is only read."""
    import os
    import shutil
    dep = Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies"
    runtime = dep / "python/bin/python3"
    renderers = sorted((Path.home() / ".codex/plugins/cache/openai-primary-runtime/documents")
                       .glob("*/skills/documents/render_docx.py"))
    if not runtime.is_file() or not renderers:
        raise SystemExit("LibreOffice printing needs the Codex document runtime (render_docx.py)")
    env = dict(os.environ)
    env["PATH"] = os.pathsep.join([str(dep / "bin/override"), str(runtime.parent), env["PATH"]])
    with tempfile.TemporaryDirectory(prefix="office-pdf-") as tmp:
        src = Path(tmp) / "in.docx"
        shutil.copyfile(docx, src)
        subprocess.run([str(runtime), str(renderers[-1]), str(src), "--output_dir", str(Path(tmp) / "out"),
                        "--emit_pdf"], env=env, check=True, capture_output=True)
        shutil.copyfile(Path(tmp) / "out" / "in.pdf", pdf)


def file_tokens(docx_path, stop_at, skip_words):
    """The original file read as draft_units reads a draft: {section id: [(word, paragraph index, start, end)]},
    a table row's cells joined by "·" (paragraph index None), guidance (all italic) left out, and the paragraphs
    an image change takes over (`skip_words`: their texts) left out. Also returns every paragraph and its text."""
    import docx as wd
    import track_changes as tc
    doc = wd.Document(docx_path)
    body = doc.element.body
    paras = tc.paragraphs(body)
    texts = [tc.para_text(p) for p in paras]
    index = {id(p): k for k, p in enumerate(paras)}
    W = WNS

    def guidance(p):
        runs = [r for r in p.iter(W + "r") if "".join(t.text or "" for t in r.iter(W + "t")).strip()]
        return bool(runs) and all(r.find(f"{W}rPr/{W}i") is not None for r in runs)

    def words(k):
        return [(m.group(), k, m.start(), m.end()) for m in re.finditer(r"\S+", texts[k])]
    out, sid = {}, None
    for el in body:
        if el.tag == W + "tbl":
            if sid is None:
                continue
            for row in el.findall(W + "tr")[1:]:
                first = True
                for p in row.iter(W + "p"):
                    k = index.get(id(p))
                    if k is None or guidance(p) or not texts[k].strip() or texts[k].strip() in skip_words:
                        continue
                    if not first:
                        out[sid].append(("·", None, 0, 0))
                    first = False
                    out[sid] += words(k)
            continue
        if el.tag != W + "p" or id(el) not in index:
            continue
        k = index[id(el)]
        clean = re.sub(r"\s+", " ", texts[k]).strip()
        if stop_at and stop_at in clean:
            break
        st = el.find(f"{W}pPr/{W}pStyle")
        if st is not None and st.get(W + "val") in ("Heading1", "Heading 1") and clean:
            m = re.match(r"(\d+)\s", clean)
            sid = m.group(1) if m else {"References": "R", "Overview Figure": "F"}.get(clean, clean)
            out[sid] = []
            continue
        if re.match(r"Task \d+(\.\d+)?:", clean):
            sid = "title"
            out[sid] = []
        if sid is None or not clean or guidance(el) or clean in skip_words:
            continue
        out[sid] += words(k)
    return out, paras, texts


def regions_to_edits(texts, stream, regions, tag):
    """Per paragraph, the smallest finds that are unique in the file and hold every changed span: regions are
    (paragraph, start, end, new text); returns [{id, find, new}]."""
    edits, by = [], {}
    stream = stream + "\n" + "\n".join(r[3] for r in regions)   # unique also among the words being added
    for k, a, b, new in regions:
        by.setdefault(k, []).append([a, b, [(a, b, new)]])
    for k, groups in by.items():
        t = texts[k]
        bounds = [(m.start(), m.end()) for m in re.finditer(r"\S+", t)]
        groups.sort(key=lambda g: g[0])
        while True:
            merged = []
            for g in groups:                         # spans that touch, or sit two words apart, become one
                if merged and len(merged[-1]) == 3 and len(t[merged[-1][1]:g[0]].split()) <= 2:
                    merged[-1][1] = max(merged[-1][1], g[1])
                    merged[-1][2] += g[2]
                else:
                    merged.append(g)
            groups, grown = merged, False
            for g in groups:
                s = t[g[0]:g[1]]
                if len(g) > 3 or (s.strip() and stream.count(s) == 1):
                    continue
                left = [x for x in bounds if x[1] <= g[0]]
                right = [x for x in bounds if x[0] >= g[1]]
                if left:                              # the words before first, then after
                    g[0] = left[-1][0]
                elif right:
                    g[1] = right[0][1]
                else:                                 # the whole paragraph repeats elsewhere ("(Refs)"): find it
                    prev = next((texts[q] for q in range(k - 1, -1, -1)   # after the nearest paragraph before it
                                 if q not in by and texts[q].strip()        # that no change touches
                                 and stream.count(texts[q].strip()) == 1), None)
                    if prev is None or not s.strip():
                        raise SystemExit(f"paragraph {k}: no unique words around a change ({t[:60]!r})")
                    g.append(prev.strip())
                    continue
                grown = True
            if not grown:
                break
        for g in groups:
            new, pos = "", g[0]
            for a, b, rep in sorted(g[2]):
                new += t[pos:a] + rep
                pos = b
            new += t[pos:g[1]]
            edits.append({"id": f"{tag}{len(edits) + 1}", "find": t[g[0]:g[1]], "new": new,
                          **({"after": [g[3]]} if len(g) > 3 else {})})
    return edits


def check_delivery(docx_path, tracked, sents, stop_at):
    """The delivered file's two checks: rejecting every change gives the original file's text, paragraph by
    paragraph, and accepting every change gives the draft's words, section by section. Returns (ok, messages)."""
    import zipfile
    from lxml import etree
    W = WNS

    def body(path):
        root = etree.fromstring(zipfile.ZipFile(path).read("word/document.xml"))
        return [p for p in root.iter(W + "p") if not any(a.tag == W + "txbxContent" for a in p.iterancestors())]

    def text(p, accept):
        out = ""
        for e in p.iter(W + "t", W + "delText"):
            gone = any(a.tag == W + ("del" if accept else "ins") for a in e.iterancestors())
            out += "" if gone else (e.text or "")
        return out
    norm = lambda ts: [re.sub(r"\s+", " ", t).strip() for t in ts if t.strip()]
    rejected, carry = [], ""
    for p in body(tracked):                             # an inserted paragraph mark joins its paragraph to the next
        t = carry + text(p, False)
        if p.find(f"{W}pPr/{W}rPr/{W}ins") is not None:
            carry = t
            continue
        carry = ""
        rejected.append(t)
    msgs, ok = [], norm([text(p, True) for p in body(docx_path)]) == norm(rejected)
    msgs.append(f"reject every change = the original file: {ok}")
    accepted, _, _ = file_tokens(tracked, stop_at, set())
    bad = [sid for sid, units in sents.items()
           if [w for u in units for w in u[2].split() if w != "·"] != [o[0] for o in accepted.get(sid, []) if o[0] != "·"]]
    msgs.append("accept every change = the draft: " + ("True" if not bad else f"False, in sections {bad}"))
    return ok and not bad, msgs


def cmd_deliver(a):
    """Write the tracked Word file from the topic's draft (JL 261011: "the delivery, we can think about it in the
    last step"): the draft's words against the original file's, word by word, become the smallest tracked changes
    by the plan's author; a draft paragraph that starts inside a file paragraph becomes a tracked new paragraph;
    the plan's image changes and its deletions outside the draft's text (guidance) are kept as they are; and each
    evidence item becomes a Word comment on its anchor words. Writes the plan's `draft` file and its PDF."""
    import difflib
    sys.path.insert(0, str(HERE))
    import track_changes as tc
    face, plan, docx_path = load(a.face, a.key)
    board = next((p for p in face.parents if (p / "board.md").is_file()), face.parent)
    if not plan.get("sentences") or not plan.get("draft"):
        raise SystemExit("deliver needs the plan's `sentences` (the draft folder) and `draft` (the Word file to write)")
    sfile = draft_file(face.parent / plan["sentences"])
    sents = read_draft(sfile, docx_path)
    edits = plan["edits"]
    import docx as wd
    paras0 = tc.paragraphs(wd.Document(docx_path).element.body)
    skip = set()                                       # what an image change takes over stays out of the diff
    for e in edits:
        if e.get("image"):
            n, _, _ = tc.locate(paras0, e)
            cell = next((x for x in paras0[n].iterancestors() if x.tag == tc.W("w:tc")), None)
            for q in (cell.iter(tc.W("w:p")) if cell is not None else [paras0[n]]):
                skip.add(tc.para_text(q).strip())
    orig, paras, texts = file_tokens(docx_path, plan.get("stop_at"), skip - {""})
    stream = "\n".join(texts)
    regions, splits = [], []
    for sid, units in sents.items():
        O = orig.get(sid, [])
        D = [(w, u[3].rsplit(".", 1)[0], n) for n, u in enumerate(units) for w in u[2].split()]
        ops = difflib.SequenceMatcher(None, [o[0] for o in O], [d[0] for d in D], autojunk=False).get_opcodes()
        for op, i1, i2, j1, j2 in ops:
            if op == "equal":
                continue
            new_words = [d for d in D[j1:j2] if d[0] != "·"]
            parts, last = [], None                    # the new words, a sentence of the References a line each
            for w, para, n in new_words:
                sep = "" if last is None else ("\n" if sid == "R" and n != last else " ")
                parts.append(sep + w)
                last = n
            new = "".join(parts)
            real = [o for o in O[i1:i2] if o[1] is not None]
            if real:
                by_para = {}
                for o in real:
                    by_para.setdefault(o[1], []).append(o)
                for m, (k, ws) in enumerate(by_para.items()):
                    s, e_ = ws[0][2], ws[-1][3]
                    rep = new if m == 0 else ""
                    if not rep:                       # a deletion takes one space with it
                        t = texts[k]
                        if e_ < len(t) and t[e_] == " ":
                            e_ = len(t[:e_]) + len(t[e_:]) - len(t[e_:].lstrip(" "))
                        elif s > 0 and t[s - 1] == " ":
                            s = len(t[:s].rstrip(" "))
                    regions.append((k, s, e_, rep))
                continue
            before = next((O[i] for i in range(i1 - 1, -1, -1) if O[i][1] is not None), None)
            after = next((O[i] for i in range(i2, len(O)) if O[i][1] is not None), None)
            at_end = before is not None and not texts[before[1]][before[3]:].strip()
            with_next = (after is not None and j2 < len(D) and D[j2 - 1][2] == D[j2][2]   # the inserted words open
                         and (before is None or before[1] != after[1]))                   # the next paragraph's sentence
            if before is not None and not with_next and (after is None or after[1] != before[1] or not at_end):
                lead = "\n" if (at_end and (sid == "R" or (j1 and D[j1][1] != D[j1 - 1][1]))) else " "
                regions.append((before[1], before[3], before[3], lead + new))
            elif after is not None:
                regions.append((after[1], after[2], after[2], new + " "))
        for n in range(1, len(units)):               # a draft paragraph that starts inside a file paragraph
            if units[n][3].rsplit(".", 1)[0] != units[n - 1][3].rsplit(".", 1)[0] and sid != "R" \
                    and " · " not in units[n][2]:              # a table row is its own row already
                words = units[n][2].split()
                k = min(len(words), 8)
                whole = "\n".join(u[2] for v in sents.values() for u in v)
                while k < len(words) and whole.count(" ".join(words[:k])) != 1:
                    k += 1
                splits.append({"id": f"p{len(splits) + 1}", "split": " ".join(words[:k])})
    tracked = regions_to_edits(texts, stream, regions, "w")
    keep = [e for e in edits if e.get("image") or (e.get("new") == "" and e["find"] not in
            "\n".join(texts[o[1]] for v in orig.values() for o in v if o[1] is not None))]
    whole = "\n".join(u[2] for v in sents.values() for u in v)
    comments = []
    for e in edits:
        if e.get("image"):
            continue                                  # an image change keeps its own comment
        for k, ev in enumerate(evidence_of(e)):
            anchor = ev.get("anchor", "")
            seen = whole.count(anchor) if anchor else 0
            if seen != 1:
                why = ("no anchor given" if not anchor else
                       f"anchor {anchor!r} is not in the draft" if seen == 0 else
                       f"anchor {anchor!r} appears {seen} times in the draft; make it unique in the plan")
                print(f"change {e['id']}: {why}")
                tagged = next((u[2] for v in sents.values() for u in v if str(e["id"]) in u[1]), "")
                words = tagged.split()
                n = min(len(words), 6)
                while n < len(words) and whole.count(" ".join(words[:n])) != 1:
                    n += 1
                anchor = " ".join(words[:n]) if words and whole.count(" ".join(words[:n])) == 1 else ""
                print(f"  so its comment sits on the first words of its sentence, {anchor!r}" if anchor else
                      "  and its sentence has no unique first words; the comment is left out")
            if anchor:
                comments.append({"id": f"c{e['id']}-{k}", "find": anchor, "evidence": dict(ev, anchor=anchor)})
    out = board / plan["draft"]
    figure = (face.parent / plan["figure"]).resolve() if plan.get("figure") else None
    done = tc.write_tracked(docx_path, tracked + keep + splits + comments, out, author=plan.get("author", "JL"),
                            figure=figure)
    print(f"wrote {out.relative_to(board)}: {len(tracked)} tracked changes from {sfile.name}, {len(splits)} paragraph "
          f"breaks, {len(keep)} kept from the plan, {done['comments']} comments")
    ok, msgs = check_delivery(docx_path, out, sents, plan.get("stop_at"))
    for m in msgs:
        print("  check: " + m)
    if not ok:
        raise SystemExit(f"{out.name} does not match the draft; it is written but not trusted, see the checks above")
    if not a.no_pdf:
        office_pdf(out, out.with_suffix(".pdf"))
        print(f"printed {out.with_suffix('.pdf').relative_to(board)} (LibreOffice, tracked changes shown)")
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    d = sub.add_parser("draw")
    d.add_argument("face")
    d.add_argument("--key", required=True)
    d.add_argument("--out")
    d.add_argument("--pdf", help="another PDF of the file, for a test; the drawing is meant for Word's own PDF")
    d.add_argument("--updated-pdf", help="another PDF of the updated draft, for a test")
    ed = sub.add_parser("edits", help="list the sentences a person changed on the drawing")
    ed.add_argument("face")
    ed.add_argument("--key", required=True)
    ed.add_argument("--drawing")
    ed.add_argument("--apply", action="store_true", help="write the edits into the sentence draft")
    sd = sub.add_parser("seed", help="write the topic's sentence draft once, from the tracked Word file")
    sd.add_argument("face")
    sd.add_argument("--key", required=True)
    sd.add_argument("--force", action="store_true")
    dv = sub.add_parser("deliver", help="write the tracked Word file from the draft, the last step")
    dv.add_argument("face")
    dv.add_argument("--key", required=True)
    dv.add_argument("--no-pdf", action="store_true")
    a = ap.parse_args()
    {"draw": cmd_draw, "edits": cmd_edits, "seed": cmd_seed, "deliver": cmd_deliver}[a.cmd](a)


if __name__ == "__main__":
    main()
