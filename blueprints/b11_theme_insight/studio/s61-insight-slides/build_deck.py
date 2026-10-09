"""s61 · build the deck: slides.py → ../../delivery/insight-slides/.

The same build as b12's s61 build_deck.py: one 1280 × 720 SVG per slide (svg/NN-<slug>.svg) drawn by b12's
slide_kit.py, merged into insight-slides.pdf, and insight-slides.pptx with native, editable shapes and text
boxes and each slide's speaker notes. The style is the html-to-svg skill's (white, Times New Roman, ink and a
navy accent, no emoji), and its scripts do the PDF and PowerPoint work. delivery/ is generated: change
slides.py, then rerun this; never edit the deck by hand.

    python build_deck.py            (needs python-pptx, cairosvg, pypdf)
"""
import html
import importlib.util
import io
import os
import sys
from pathlib import Path

try:
    import cairosvg  # noqa: F401
except OSError:                                       # macOS: Homebrew's cairo is off the loader's path
    if os.environ.get("DYLD_FALLBACK_LIBRARY_PATH") != "/opt/homebrew/lib" and Path("/opt/homebrew/lib").exists():
        os.execve(sys.executable, [sys.executable, *sys.argv],
                  {**os.environ, "DYLD_FALLBACK_LIBRARY_PATH": "/opt/homebrew/lib"})
    raise

HERE = Path(__file__).resolve().parent
OUT = HERE.parents[1] / "delivery" / "insight-slides"
SKILL = HERE.parents[3] / "plugins" / "haipipe-toolkit" / "skills" / "1_base" / "display" / "html-to-svg" / "scripts"
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(SKILL))
import svg_deck as D  # noqa: E402  (the skill's palette and font)
from build_s61_insight_slides import K  # noqa: E402  (b12's slide_kit, loaded by path)
from slides import SLIDES  # noqa: E402

COLOR = {"ink": D.INK, "sub": D.SUB, "mut": D.MUT, "acc": D.ACC, "border": D.BORDER, "warn": D.WARN,
         "fill2": D.SURF2}


class SvgPen:
    """slide_kit's pen, as SVG (b12's build_deck.py SvgPen)."""
    cw = 0.47                                         # Times New Roman

    def __init__(self):
        self.el = []

    def rect(self, x, y, w, h, role="border", dashed=False, sw=1, fill=None):
        dash = ' stroke-dasharray="6,4"' if dashed else ""
        self.el.append(f'<rect x="{x:.1f}" y="{y:.1f}" width="{w:.1f}" height="{h:.1f}" rx="3" '
                       f'fill="{COLOR[fill] if fill else "none"}" stroke="{COLOR[role]}" stroke-width="{sw}"{dash}/>')

    def line(self, x1, y1, x2, y2, role="border"):
        self.el.append(f'<line x1="{x1:.1f}" y1="{y1:.1f}" x2="{x2:.1f}" y2="{y2:.1f}" '
                       f'stroke="{COLOR[role]}" stroke-width="1"/>')

    def text(self, x, y, s, fs, role="ink", bold=False, anchor="start"):
        self.el.append(f'<text x="{x:.1f}" y="{y + fs * 0.82:.1f}" font-size="{fs}" fill="{COLOR[role]}" '
                       f'font-weight="{"bold" if bold else "normal"}" font-family="{D.FONT}" '
                       f'text-anchor="{anchor}">{html.escape(s)}</text>')

    def arrow(self, pts, role="ink"):
        c = COLOR[role]
        self.el.append(f'<polyline points="{" ".join(f"{a:.1f},{b:.1f}" for a, b in pts)}" fill="none" '
                       f'stroke="{c}" stroke-width="1.6"/>')
        (x1, y1), (x2, y2) = pts[-2], pts[-1]
        dx, dy = (x2 > x1) - (x2 < x1), (y2 > y1) - (y2 < y1)        # the last leg runs straight
        px, py = -dy, dx
        head = [(x2 - 10 * dx + 6 * px, y2 - 10 * dy + 6 * py), (x2, y2), (x2 - 10 * dx - 6 * px, y2 - 10 * dy - 6 * py)]
        self.el.append(f'<polyline points="{" ".join(f"{a:.1f},{b:.1f}" for a, b in head)}" fill="none" '
                       f'stroke="{c}" stroke-width="1.6"/>')

    def svg(self, title):
        return (f'<svg xmlns="http://www.w3.org/2000/svg" width="{K.W}" height="{K.H}" viewBox="0 0 {K.W} {K.H}">\n'
                f'<title>{html.escape(title)}</title>\n<rect width="{K.W}" height="{K.H}" fill="#ffffff"/>\n'
                + "\n".join(self.el) + "\n</svg>\n")


def main():
    svg_dir = OUT / "svg"
    svg_dir.mkdir(parents=True, exist_ok=True)
    for old in svg_dir.glob("[0-9][0-9]-*.svg"):
        old.unlink()
    paths = []
    for s in SLIDES:
        pen = SvgPen()
        K.draw_slide(pen, s)
        p = svg_dir / f"{s['n']:02d}-{s['slug']}.svg"
        p.write_text(pen.svg(s["title"]), encoding="utf-8")
        paths.append(p)

    import cairosvg
    from pypdf import PdfReader, PdfWriter
    pdf = PdfWriter()
    for p in paths:
        pdf.append(PdfReader(io.BytesIO(cairosvg.svg2pdf(url=str(p)))))
    with open(OUT / "insight-slides.pdf", "wb") as f:
        pdf.write(f)

    import build_pptx_native as N                     # the skill's SVG → native-shapes PowerPoint
    from pptx import Presentation
    from pptx.util import Emu
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(K.W * N.PX), Emu(K.H * N.PX)
    for s, p in zip(SLIDES, paths):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        N.build_slide(slide.shapes, p)
        slide.notes_slide.notes_text_frame.text = s["notes"]
    prs.save(OUT / "insight-slides.pptx")
    print(f"wrote {OUT.relative_to(HERE.parents[3].parent)}: {len(paths)} slides (svg/, insight-slides.pdf, "
          f"insight-slides.pptx)")


if __name__ == "__main__":
    main()
